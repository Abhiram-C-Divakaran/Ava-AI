# main.py
import os
import time
import uuid
import json
import shutil
import zipfile
import io
import re
import secrets
import aiohttp
import logging
from datetime import datetime, timezone
from contextlib import asynccontextmanager
from fastapi import FastAPI, HTTPException, Request, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse, Response
from starlette.middleware.sessions import SessionMiddleware
from pydantic import BaseModel, EmailStr, Field
from typing import Optional, Any, Union
from io import BytesIO

from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

from llm import call_llm, call_llm_streaming, call_llm_with_constraints
from MLpipeline import SentimentAnalyzer, IntentClassifier, FrustrationDetector
import database as db
import search
import memory
import adaptation
import powers
import config
from version import __version__
from metrics import metrics

ENVIRONMENT = config.ENVIRONMENT
SECRET_KEY = config.SECRET_KEY
if not SECRET_KEY:
    if ENVIRONMENT == "production":
        raise RuntimeError("SECRET_KEY must be configured in production environment")
    SECRET_KEY = "dev_insecure_secret_key_change_in_production"

# Structured Logging
try:
    from pythonjsonlogger.json import JsonFormatter
except ImportError:
    try:
        from pythonjsonlogger import jsonlogger as JsonFormatter
    except ImportError:
        JsonFormatter = None

logger = logging.getLogger("ava")
logger.setLevel(logging.INFO)
if not logger.handlers:
    sh = logging.StreamHandler()
    if JsonFormatter:
        sh.setFormatter(JsonFormatter("%(asctime)s %(levelname)s %(name)s %(message)s"))
    else:
        sh.setFormatter(logging.Formatter("[%(asctime)s] %(levelname)s: %(message)s"))
    logger.addHandler(sh)

def log_event(event_name: str, **kwargs):
    safe_data = {"event": event_name}
    for k, v in kwargs.items():
        if k in ("password", "token", "access_token", "secret", "cookie", "code", "authorization"):
            continue
        safe_data[k] = v
    logger.info(json.dumps(safe_data))

def rate_limit_key(request: Request) -> str:
    user_id = request.session.get("user_id")
    if user_id:
        return f"user:{user_id}"
    return f"ip:{get_remote_address(request)}"

limiter = Limiter(key_func=rate_limit_key)
if config.ENVIRONMENT != "production" and os.environ.get("DISABLE_RATE_LIMIT", "").lower() in ("true", "1"):
    limiter.enabled = False

# OAuth
from authlib.integrations.starlette_client import OAuth, OAuthError

try:
    from pptx import Presentation
    PPT_AVAILABLE = True
except ImportError:
    PPT_AVAILABLE = False

try:
    import replicate
    REPLICATE_AVAILABLE = True
except ImportError:
    REPLICATE_AVAILABLE = False

# ─── Lifespan ──────────────────────────────────────────────────────────────
@asynccontextmanager
async def lifespan(app: FastAPI):
    val = config.validate_config()
    if not val["valid"]:
        logger.error(f"Configuration validation: {val['errors']}")
        if config.ENVIRONMENT == "production":
            raise RuntimeError(f"Invalid production configuration: {val['errors']}")
    config.ensure_directories()
    db.init_db()
    yield

app = FastAPI(title="Ava AI", version=__version__, lifespan=lifespan)
app.state.limiter = limiter

def custom_rate_limit_handler(request: Request, exc: RateLimitExceeded):
    metrics.inc("rate_limit_events")
    return _rate_limit_exceeded_handler(request, exc)

app.add_exception_handler(RateLimitExceeded, custom_rate_limit_handler)

ALLOWED_ORIGINS = config.ALLOWED_ORIGINS

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.add_middleware(
    SessionMiddleware,
    secret_key=SECRET_KEY,
    same_site="lax",
    https_only=(ENVIRONMENT == "production"),
)

@app.middleware("http")
async def correlation_id_middleware(request: Request, call_next):
    request_id = request.headers.get("X-Request-ID") or str(uuid.uuid4())
    request.state.request_id = request_id
    response = await call_next(request)
    response.headers["X-Request-ID"] = request_id
    return response

# ─── Health & Readiness Endpoints ──────────────────────────────────────────
@app.get("/health")
def health_check():
    return {"status": "ok", "version": __version__}

@app.get("/ready")
def ready_check():
    errors = []
    val = config.validate_config()
    if not val["valid"]:
        errors.extend(val["errors"])

    try:
        with db.get_conn() as conn:
            conn.execute("CREATE TABLE IF NOT EXISTS _readiness_check (id INTEGER PRIMARY KEY, ts REAL)")
            conn.execute("INSERT OR REPLACE INTO _readiness_check (id, ts) VALUES (1, ?)", (time.time(),))
            conn.commit()
            conn.execute("SELECT ts FROM _readiness_check WHERE id = 1").fetchone()
    except Exception as e:
        metrics.inc("database_errors")
        logger.exception("Readiness check: DB failed")
        errors.append(f"database: {str(e)}")

    for dir_path, label in [(config.UPLOAD_DIR, "uploads"), (config.REVIEW_IMAGE_DIR, "review_images")]:
        try:
            os.makedirs(dir_path, exist_ok=True)
            test_file = os.path.join(dir_path, ".write_test")
            with open(test_file, "w") as f:
                f.write("ok")
            if os.path.exists(test_file):
                os.remove(test_file)
        except Exception as e:
            logger.exception(f"Readiness check: Storage {label} not writable")
            errors.append(f"storage_{label}: not writable ({str(e)})")

    if errors:
        return Response(
            content=json.dumps({"status": "unhealthy", "errors": errors, "version": __version__}),
            status_code=503,
            media_type="application/json"
        )

    return {
        "status": "ready",
        "version": __version__,
        "database": "connected",
        "storage": "writable"
    }

@app.get("/api/metrics")
def get_operational_metrics(request: Request):
    require_admin(request)
    return metrics.get_snapshot()

ADMIN_PASSWORD = config.ADMIN_PASSWORD

def require_admin(request: Request):
    if not request.session.get("is_admin"):
        raise HTTPException(status_code=401, detail="Admin authentication required")

def require_user(request: Request) -> str:
    user_id = request.session.get("user_id")
    if not user_id:
        raise HTTPException(status_code=401, detail="Authentication required")
    return user_id

def require_same_user(request: Request, requested_user_id: str) -> str:
    current_user = require_user(request)
    if current_user != requested_user_id and not request.session.get("is_admin"):
        raise HTTPException(status_code=403, detail="Forbidden: access denied")
    return current_user

oauth = OAuth()
oauth.register(
    name='google',
    client_id=os.getenv("GOOGLE_CLIENT_ID"),
    client_secret=os.getenv("GOOGLE_CLIENT_SECRET"),
    server_metadata_url='https://accounts.google.com/.well-known/openid-configuration',
    client_kwargs={'scope': 'openid email profile'},
    redirect_uri=os.getenv("GOOGLE_REDIRECT_URI", "http://localhost:8000/api/auth/google/callback")
)

sentiment_analyzer = SentimentAnalyzer()
intent_classifier = IntentClassifier()
frustration_detector = FrustrationDetector()

def is_support_issue(message: str, intent: str) -> bool:
    support_intents = ["billing_issue", "login_issue", "app_crash", "performance_issue",
                       "data_loss", "connectivity", "installation", "escalation_request"]
    if intent in support_intents:
        return True
    support_keywords = ["error", "crash", "not working", "broken", "help", "issue",
                        "problem", "can't", "doesn't", "failed", "login", "password",
                        "refund", "payment"]
    msg_lower = message.lower()
    return any(kw in msg_lower for kw in support_keywords)

# Language name lookup for common codes
_LANG_NAMES = {
    "en": "English", "es": "Spanish", "fr": "French", "de": "German",
    "it": "Italian", "pt": "Portuguese", "ru": "Russian", "zh": "Chinese",
    "ja": "Japanese", "ko": "Korean", "ar": "Arabic", "hi": "Hindi",
    "nl": "Dutch", "pl": "Polish", "tr": "Turkish", "sv": "Swedish",
    "da": "Danish", "fi": "Finnish", "no": "Norwegian", "uk": "Ukrainian",
    "vi": "Vietnamese", "th": "Thai", "id": "Indonesian", "ms": "Malay",
    "ro": "Romanian", "cs": "Czech", "hu": "Hungarian", "el": "Greek",
}

def _detect_translation(message: str) -> dict | None:
    msg = message.strip().lower()
    m = re.search(
        r"\b(translate|convert|render)\b.*?\bto\b\s+([a-z]+)", msg, re.IGNORECASE
    )
    if m:
        tgt_word = m.group(2).strip().title()
        src_m = re.search(r"\bfrom\b\s+([a-z]+)", msg, re.IGNORECASE)
        src_word = src_m.group(1).strip().title() if src_m else "auto-detect"
        return {"src": src_word, "tgt": tgt_word}
    m2 = re.search(
        r"\b(how do you say|what is|what'?s|say it|write it)\b.*?\bin\b\s+([a-z]+)", msg, re.IGNORECASE
    )
    if m2:
        tgt_word = m2.group(2).strip().title()
        return {"src": "auto-detect", "tgt": tgt_word}
    return None

# ─── Pydantic models ────────────────────────────────────────────────────
class ChatRequest(BaseModel):
    user_id: Optional[str] = None
    message: str = Field(..., min_length=1, max_length=20000)
    session_id: Optional[str] = None
    doc_id: Optional[str] = None
    image_base64: Optional[str] = None
    image_mime: Optional[str] = "image/jpeg"
    file_context: Optional[str] = None
    file_name: Optional[str] = None
    mode: Optional[str] = "flash"
    web_search: Optional[bool] = False

class FeedbackRequest(BaseModel):
    user_id: Optional[str] = None
    session_id: str
    message_id: str
    helpful: bool

class SignupRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=100)
    email: EmailStr
    password: str = Field(..., min_length=8, max_length=128)
    terms_accepted: Optional[bool] = False

class LoginRequest(BaseModel):
    email: EmailStr
    password: str = Field(..., min_length=1, max_length=128)

class SuggestionRequest(BaseModel):
    user_id: Optional[str] = None
    current_input: str = Field(..., max_length=1000)

class PPTRequest(BaseModel):
    user_id: Optional[str] = None
    topic: str = Field(..., min_length=1, max_length=200)
    slides: int = Field(default=5, ge=1, le=20)

class ImageRequest(BaseModel):
    user_id: Optional[str] = None
    prompt: str = Field(..., min_length=1, max_length=1000)

class PreferencesUpdate(BaseModel):
    custom_instructions: Optional[str] = Field(None, max_length=5000)
    personality: Optional[str] = Field(None, max_length=50)
    theme: Optional[str] = Field(None, max_length=20)
    birth_date: Optional[str] = Field(None, max_length=20)
    gender: Optional[str] = Field(None, max_length=20)

class RenameRequest(BaseModel):
    new_title: str = Field(..., min_length=1, max_length=255)

class ChangePasswordRequest(BaseModel):
    current_password: str
    new_password: str = Field(..., min_length=8, max_length=128)

class TermsAcceptRequest(BaseModel):
    terms_accepted: bool

# ─── Auth ──────────────────────────────────────────────────────────────
@app.post("/api/auth/signup")
@limiter.limit("5/minute")
def signup(req: SignupRequest, request: Request):
    if len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters")
    if db.email_exists(req.email):
        raise HTTPException(status_code=409, detail="Email already exists")
    
    user_id = str(uuid.uuid4())
    db.create_user(
        user_id=user_id, 
        name=req.name.strip(), 
        email=req.email, 
        password=req.password,
        terms_accepted=req.terms_accepted
    )
    request.session["user_id"] = user_id
    log_event("user_signup_success")
    return {
        "user_id": user_id, 
        "name": req.name.strip(), 
        "email": req.email.lower(),
        "terms_accepted": req.terms_accepted,
        "is_admin": False
    }

@app.post("/api/auth/login")
@limiter.limit("5/minute")
def login(req: LoginRequest, request: Request):
    user = db.get_user_by_email(req.email)
    if not user or not db.verify_password(req.password, user["password_hash"], user["password_salt"]):
        log_event("login_failed")
        raise HTTPException(status_code=401, detail="Invalid email or password")
    request.session["user_id"] = user["user_id"]
    is_admin = user.get("is_admin", 0) == 1
    if is_admin:
        request.session["is_admin"] = True
        request.session["admin_verified"] = True
    log_event("login_success", is_admin=is_admin)
    return {
        "user_id": user["user_id"],
        "name": user["name"],
        "email": user["email"],
        "is_admin": is_admin,
        "terms_accepted": bool(user.get("terms_accepted", 0))
    }

@app.post("/api/auth/logout")
async def logout(request: Request):
    request.session.clear()
    log_event("logout_success")
    return {"status": "logged out"}

# ─── Google OAuth ──────────────────────────────────────────────────────
@app.get("/api/auth/google")
async def google_login(request: Request):
    redirect_uri = os.getenv("GOOGLE_REDIRECT_URI", "http://localhost:8000/api/auth/google/callback")
    return await oauth.google.authorize_redirect(request, redirect_uri)

@app.get("/api/auth/google/callback")
async def google_auth(request: Request):
    """
    Google OAuth callback handler.
    Uses userinfo endpoint as fallback when id_token is not available.
    """
    try:
        token = await oauth.google.authorize_access_token(request)
        user_info = None
        
        if token and 'id_token' in token:
            try:
                user_info = await oauth.google.parse_id_token(request, token)
            except Exception:
                user_info = None
        
        if not user_info and token:
            access_token = token.get('access_token')
            if not access_token:
                raise HTTPException(status_code=400, detail="No access token received")
            
            timeout = aiohttp.ClientTimeout(total=10)
            async with aiohttp.ClientSession(timeout=timeout) as session:
                async with session.get(
                    'https://openidconnect.googleapis.com/v1/userinfo',
                    headers={'Authorization': f'Bearer {access_token}'}
                ) as resp:
                    if resp.status != 200:
                        raise HTTPException(status_code=400, detail="Failed to retrieve user info")
                    user_info = await resp.json()
        
        if not user_info:
            raise HTTPException(status_code=400, detail="Failed to get user information")
        
        email = user_info.get('email')
        if not email:
            raise HTTPException(status_code=400, detail="No email provided by Google")
        
        name = user_info.get('name', email.split('@')[0])
        if not name or name.strip() == '':
            name = email.split('@')[0]
        
        user = db.get_user_by_email(email)
        if not user:
            user_id = str(uuid.uuid4())
            temp_pw = secrets.token_urlsafe(16)
            db.create_user(
                user_id=user_id, 
                name=name, 
                email=email, 
                password=temp_pw, 
                terms_accepted=True
            )
        else:
            user_id = user['user_id']
        
        request.session["user_id"] = user_id
        if user and user.get("is_admin", 0) == 1:
            request.session["is_admin"] = True
            request.session["admin_verified"] = True
        log_event("oauth_login_success", provider="google")
        
        safe_name = name.replace("'", "\\'")
        html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>Authentication Complete</title>
            <style>
                body {{
                    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                    background: #0F0A2A;
                    color: #F0EEFF;
                    display: flex;
                    align-items: center;
                    justify-content: center;
                    height: 100vh;
                    margin: 0;
                    padding: 20px;
                }}
                .container {{
                    text-align: center;
                    padding: 40px;
                    background: rgba(255,255,255,0.05);
                    border-radius: 16px;
                    border: 1px solid rgba(255,255,255,0.1);
                    max-width: 400px;
                    width: 100%;
                }}
                .spinner {{
                    width: 40px;
                    height: 40px;
                    border: 3px solid rgba(255,255,255,0.1);
                    border-top-color: #54E6F2;
                    border-radius: 50%;
                    animation: spin 0.8s linear infinite;
                    margin: 20px auto;
                }}
                @keyframes spin {{
                    to {{ transform: rotate(360deg); }}
                }}
                .success {{
                    color: #4ADE80;
                    font-size: 2.5rem;
                    margin-bottom: 10px;
                }}
                h2 {{
                    font-weight: 600;
                    margin: 0 0 8px 0;
                }}
                p {{
                    color: #B8B0D8;
                    margin: 8px 0;
                }}
                .sub {{
                    color: #7A6F9A;
                    font-size: 0.8rem;
                    margin-top: 20px;
                }}
            </style>
        </head>
        <body>
            <div class="container">
                <div class="success">✅</div>
                <h2>Authentication Successful!</h2>
                <p>Welcome, {name}!</p>
                <div class="spinner"></div>
                <p>Redirecting to chat...</p>
                <p class="sub">You will be redirected automatically.</p>
            </div>
            <script>
                // Store user data in localStorage
                const userData = {{
                    user_id: '{user_id}',
                    name: '{safe_name}',
                    email: '{email}',
                    terms_accepted: true,
                    is_oauth: true
                }};
                localStorage.setItem('neurosupport_user', JSON.stringify(userData));
                localStorage.setItem('ava-terms-accepted', 'true');
                localStorage.setItem('ava-terms-accepted-date', new Date().toISOString());
                
                // Redirect after a short delay
                setTimeout(() => {{
                    window.location.href = '/chat.html';
                }}, 1500);
            </script>
        </body>
        </html>
        """
        return HTMLResponse(content=html)
        
    except HTTPException:
        log_event("oauth_login_failed")
        raise
    except Exception as e:
        logger.exception("Unexpected error in google_auth")
        log_event("oauth_login_failed")
        raise HTTPException(status_code=400, detail="Authentication failed")

# ─── Admin auth ──────────────────────────────────────────────────────
@app.post("/api/auth/admin-verify")
@limiter.limit("5/minute")
async def admin_verify(request: Request, payload: dict):
    password = payload.get("admin_password")
    if not password:
        raise HTTPException(status_code=400, detail="Missing admin_password")
    if not ADMIN_PASSWORD:
        raise HTTPException(status_code=503, detail="Admin password not configured")
    if password != ADMIN_PASSWORD:
        log_event("admin_verify_failed")
        raise HTTPException(status_code=401, detail="Invalid admin password")
    request.session["is_admin"] = True
    request.session["admin_verified"] = True
    log_event("admin_verify_success")
    return {"status": "ok"}

@app.get("/api/auth/admin-status")
async def admin_status(request: Request):
    return {
        "is_admin": bool(request.session.get("is_admin")),
        "admin_verified": bool(request.session.get("admin_verified"))
    }

@app.post("/api/auth/admin-logout")
async def admin_logout(request: Request):
    request.session.pop("is_admin", None)
    request.session.pop("admin_verified", None)
    log_event("admin_logout")
    return {"status": "ok"}

# ─── Image vision via Groq ─────────────────────────────────────────────────────
def call_llm_with_vision(user_message: str, image_base64: str, image_mime: str = "image/jpeg",
                          system_prompt: str = "", file_context: str = "") -> str:
    import os as _os, requests as _req
    api_key = _os.getenv("GROQ_API_KEY", "")
    if not api_key:
        return "[Vision requires GROQ_API_KEY to be set]"
    content_parts = []
    if file_context:
        content_parts.append({"type": "text", "text": f"[Attached document:]\n{file_context[:6000]}\n\n"})
    content_parts.append({"type": "image_url", "image_url": {"url": f"data:{image_mime};base64,{image_base64}"}})
    content_parts.append({"type": "text", "text": user_message or "Describe and analyse this image in detail."})
    messages = []
    if system_prompt:
        messages.append({"role": "system", "content": system_prompt})
    messages.append({"role": "user", "content": content_parts})
    payload = {"model": "meta-llama/llama-4-scout-17b-16e-instruct", "messages": messages, "max_tokens": 1500, "temperature": 0.5}
    try:
        r = _req.post("https://api.groq.com/openai/v1/chat/completions",
                      headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                      json=payload, timeout=45)
        if r.status_code == 200:
            return r.json()["choices"][0]["message"]["content"].strip()
        return f"[Vision API error {r.status_code}: {r.text[:200]}]"
    except Exception as e:
        return f"[Vision error: {str(e)}]"

def assemble_chat_prompt_context(
    user_id: str,
    message: str,
    session_id: Optional[str] = None,
    doc_id: Optional[str] = None,
    file_context: Optional[str] = None,
    mode: Optional[str] = "flash",
    web_search: Optional[bool] = False,
    custom_instructions: str = "",
    personality: str = "friendly",
) -> tuple[str, dict]:
    """
    Constructs the augmented system prompt with correct conceptual separation:
    - user_memory_context = persistent cross-session factual memory
    - adaptation_context = learned behavioral response preferences
    - conversation_context = current-session context only (strictly session_context)
    Returns (system_prompt, aug_meta).
    """
    # 1. Session history is strictly scoped to the active chat
    session_context = memory.get_session_context(user_id, session_id) if session_id else ""
    # 2. Durable user facts stored separately
    user_memory_context = memory.get_user_memory_context(user_id)
    # Conceptual separation: conversation_context is strictly session_context
    conversation_context = session_context

    sentiment = sentiment_analyzer.analyze(message)
    intent = intent_classifier.classify(message)
    frustration = frustration_detector.score(user_id, message)

    failed = db.get_failed_solutions(user_id)
    failed_note = "Previously attempted solutions that didn't work: " + "; ".join(failed) if failed else ""

    translation_langs = _detect_translation(message)
    support_flag = is_support_issue(message, intent)

    fc = file_context or ""
    if doc_id and not fc:
        try:
            doc = db.get_document(doc_id)
            if doc:
                fc = doc.get("extracted_text", "")[:8000]
        except Exception:
            pass

    adaptation_context = ""
    policy = {}
    try:
        adaptation_context = adaptation.get_adaptation_context(user_id, message)
        policy = adaptation.get_behavior_policy(user_id, message)
    except Exception as e:
        print(f"⚠️ Adaptation context build failed: {e}")

    from llm import build_system_prompt as _bsp
    _base = _bsp(custom_instructions, personality, mode or "flash")

    system_override, aug_meta = powers.build_augmented_system_prompt(
        base_prompt=_base,
        user_message=message,
        conversation_context=conversation_context,
        file_context=fc,
        failed_note=failed_note,
        sentiment_label=sentiment["label"],
        frustration=frustration,
        translation_langs=translation_langs,
        support_flag=support_flag,
        web_search_enabled=bool(web_search),
        adaptation_context=adaptation_context,
        user_memory_context=user_memory_context,
    )
    aug_meta["sentiment"] = sentiment
    aug_meta["intent"] = intent
    aug_meta["frustration"] = frustration
    aug_meta["translation_langs"] = translation_langs
    aug_meta["support_flag"] = support_flag
    aug_meta["file_context"] = fc
    aug_meta["adaptation_context"] = adaptation_context
    aug_meta["user_memory_context"] = user_memory_context
    aug_meta["conversation_context"] = conversation_context
    aug_meta["strategy"] = policy.get("preferred_strategy")
    aug_meta["adaptation_used"] = adaptation.is_adaptation_used(user_id, message)
    aug_meta["policy"] = policy
    return system_override, aug_meta

# ─── Chat (non‑streaming) ──────────────────────────────────────────────
@app.post("/api/chat")
@limiter.limit("30/minute")
def chat(req: ChatRequest, request: Request):
    metrics.inc("chat_requests")
    current_user = require_user(request)
    if req.user_id and req.user_id != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")
    user_id = current_user

    if req.session_id:
        sess = db.get_session(req.session_id)
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found")
        if sess["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    if req.doc_id:
        doc = db.get_document(req.doc_id)
        if doc and doc["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: document does not belong to you")

    # ─── Memory storage command ──────────────────────────────────────────────
    memory_code_match = re.match(r"^Remember this code:\s*(.+)$", req.message, re.IGNORECASE)
    if memory_code_match:
        code = memory_code_match.group(1).strip()
        response_text = f"Got it. I'll remember the code {code} for this conversation."
        message_id = str(uuid.uuid4())
        session_id = req.session_id or str(uuid.uuid4())
        if not req.session_id:
            db.create_session(session_id, user_id, title=req.message[:50])
        db.save_message(
            message_id=message_id,
            session_id=session_id,
            user_id=user_id,
            user_message=req.message,
            agent_response=response_text,
            intent="memory_storage",
            sentiment={"label": "neutral", "score": 0.5},
            frustration=0.0,
            latency_ms=0,
        )
        memory.maybe_refresh_user_memory(user_id, session_id, req.message, response_text)
        return {
            "response": response_text,
            "message_id": message_id,
            "session_id": session_id,
            "metadata": {"intent": "memory_storage", "sentiment": {"label": "neutral"}, "frustration_score": 0.0}
        }

    prefs = db.get_user_preferences(user_id)
    custom_instructions = prefs.get("custom_instructions", "")
    personality = prefs.get("personality", "friendly")

    system_override, aug_meta = assemble_chat_prompt_context(
        user_id=user_id,
        message=req.message,
        session_id=req.session_id,
        doc_id=req.doc_id,
        file_context=req.file_context,
        mode=req.mode or "flash",
        web_search=getattr(req, "web_search", False),
        custom_instructions=custom_instructions,
        personality=personality,
    )
    sentiment = aug_meta["sentiment"]
    intent = aug_meta["intent"]
    frustration = aug_meta["frustration"]
    translation_langs = aug_meta["translation_langs"]
    support_flag = aug_meta["support_flag"]
    file_context = aug_meta["file_context"]

    # DeepL translation path
    if translation_langs and translation_langs.get("tgt"):
        text_to_translate = powers.extract_text_to_translate(req.message)
        if text_to_translate:
            deepl_result = powers.translate_deepl(
                text_to_translate,
                target_lang_name=translation_langs["tgt"],
                source_lang_name=translation_langs.get("src", "auto"),
            )
            if deepl_result:
                response = f"{deepl_result['translated']}\n\n*Translated from {deepl_result['source_lang']} → {deepl_result['target_lang']} via DeepL*"
                aug_meta["deepl_used"] = True
                message_id = str(uuid.uuid4())
                session_id = req.session_id or str(uuid.uuid4())
                if not req.session_id:
                    db.create_session(session_id, user_id, title=req.message[:50])
                db.save_message(message_id=message_id, session_id=session_id, user_id=user_id,
                    user_message=req.message, agent_response=response, intent=intent,
                    sentiment=sentiment, frustration=frustration, latency_ms=0)
                memory.maybe_refresh_user_memory(user_id, session_id, req.message, response)
                return {"response": response, "message_id": message_id, "session_id": session_id,
                    "metadata": {"intent": intent, "sentiment": sentiment, "frustration_score": frustration}}

    start_time = time.time()
    try:
        if req.image_base64:
            response = call_llm_with_vision(
                req.message, req.image_base64, req.image_mime or "image/jpeg",
                system_prompt=system_override, file_context=file_context
            )
        else:
            response = call_llm_with_constraints(
                req.message,
                base_system_prompt=system_override,
                custom_instructions=custom_instructions,
                personality=personality,
                max_tokens=1200,
                mode=req.mode or "flash",
            )
    except Exception as e:
        metrics.inc("llm_failures")
        metrics.inc("chat_errors")
        logger.exception("Chat LLM execution failed")
        response = "I'm sorry, I encountered an error processing your request."

    latency_ms = int((time.time() - start_time) * 1000)
    metrics.record_latency(latency_ms)
    if aug_meta.get("adaptation_used", False):
        metrics.inc("adaptation_used_count")
    message_id = str(uuid.uuid4())
    session_id = req.session_id or str(uuid.uuid4())
    if not req.session_id:
        db.create_session(session_id, user_id, title=req.message[:50])
    db.save_message(message_id=message_id, session_id=session_id, user_id=user_id,
        user_message=req.message, agent_response=response, intent=intent,
        sentiment=sentiment, frustration=frustration, latency_ms=latency_ms)
    memory.maybe_refresh_user_memory(user_id, session_id, req.message, response)
    try:
        adaptation.observe_interaction(
            user_id=user_id,
            user_message=req.message,
            agent_response=response,
            intent=intent,
            sentiment=sentiment.get("label") if isinstance(sentiment, dict) else str(sentiment),
            adaptation_used=aug_meta.get("adaptation_used", False),
        )
    except Exception as e:
        logger.warning(f"Adaptation observation failed: {e}")

    log_event(
        "chat_completed",
        latency_ms=latency_ms,
        intent=intent,
        adaptation_used=aug_meta.get("adaptation_used", False),
        strategy=aug_meta.get("strategy")
    )

    return {"response": response, "message_id": message_id, "session_id": session_id,
        "metadata": {"intent": intent, "sentiment": sentiment, "frustration_score": frustration,
                     "searched": aug_meta.get("searched", False),
                     "strategy": aug_meta.get("strategy"),
                     "adaptation_used": aug_meta.get("adaptation_used", False)}}

# ─── Chat streaming ─────────────────────────────────────────────────────
@app.post("/api/chat/stream")
@limiter.limit("30/minute")
def chat_stream(req: ChatRequest, request: Request):
    metrics.inc("stream_requests")
    current_user = require_user(request)
    if req.user_id and req.user_id != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")
    user_id = current_user

    if req.session_id:
        sess = db.get_session(req.session_id)
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found")
        if sess["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    if req.doc_id:
        doc = db.get_document(req.doc_id)
        if doc and doc["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: document does not belong to you")

    def generate():
        # ─── Memory storage command ──────────────────────────────────────────────
        memory_code_match = re.match(r"^Remember this code:\s*(.+)$", req.message, re.IGNORECASE)
        if memory_code_match:
            code = memory_code_match.group(1).strip()
            response_text = f"Got it. I'll remember the code {code} for this conversation."
            yield f"data: {json.dumps({'chunk': response_text})}\n\n"
            latency_ms = 0
            message_id = str(uuid.uuid4())
            session_id = req.session_id or str(uuid.uuid4())
            if not req.session_id:
                db.create_session(session_id, user_id, title=req.message[:50])
            db.save_message(
                message_id=message_id,
                session_id=session_id,
                user_id=user_id,
                user_message=req.message,
                agent_response=response_text,
                intent="memory_storage",
                sentiment={"label": "neutral", "score": 0.5},
                frustration=0.0,
                latency_ms=latency_ms,
            )
            memory.maybe_refresh_user_memory(user_id, session_id, req.message, response_text)
            yield f"data: {json.dumps({'done': True, 'message_id': message_id, 'session_id': session_id, 'metadata': {'intent': 'memory_storage', 'sentiment': {'label': 'neutral'}, 'frustration_score': 0.0}})}\n\n"
            return

        prefs = db.get_user_preferences(user_id)
        custom_instructions = prefs.get("custom_instructions", "")
        personality = prefs.get("personality", "friendly")

        system_override, aug_meta = assemble_chat_prompt_context(
            user_id=user_id,
            message=req.message,
            session_id=req.session_id,
            doc_id=req.doc_id,
            file_context=req.file_context,
            mode=req.mode or "flash",
            web_search=getattr(req, "web_search", False),
            custom_instructions=custom_instructions,
            personality=personality,
        )
        sentiment = aug_meta["sentiment"]
        intent = aug_meta["intent"]
        frustration = aug_meta["frustration"]
        translation_langs = aug_meta["translation_langs"]
        support_flag = aug_meta["support_flag"]
        file_context = aug_meta["file_context"]

        if aug_meta.get("searched"):
            yield f"data: {json.dumps({'status': '🔍 Searching the web…', 'searching': True})}\n\n"

        start_time = time.time()
        full_response = ""

        run_cmd = powers.detect_run_command(req.message)
        if run_cmd:
            _run_lang = run_cmd["lang"]
            _run_status = json.dumps({"status": f"⚙️ Running {_run_lang} code…"})
            yield f"data: {_run_status}\n\n"
            result = powers.execute_code(run_cmd["lang"], run_cmd["code"])
            full_response = powers.format_execution_result(run_cmd["lang"], run_cmd["code"], result)
            yield f"data: {json.dumps({'chunk': full_response})}\n\n"

        elif translation_langs and translation_langs.get("tgt"):
            text_to_translate = powers.extract_text_to_translate(req.message)
            if text_to_translate:
                yield f"data: {json.dumps({'status': '🌐 Translating…'})}\n\n"
                deepl_result = powers.translate_deepl(
                    text_to_translate,
                    target_lang_name=translation_langs["tgt"],
                    source_lang_name=translation_langs.get("src", "auto"),
                )
                if deepl_result:
                    full_response = f"{deepl_result['translated']}\n\n*Translated from {deepl_result['source_lang']} → {deepl_result['target_lang']} via DeepL*"
                    aug_meta["deepl_used"] = True
                    yield f"data: {json.dumps({'chunk': full_response})}\n\n"
                else:
                    pass

            if not full_response:
                try:
                    stream = call_llm_streaming(
                        req.message,
                        custom_instructions=custom_instructions,
                        personality=personality,
                        system_prompt_override=system_override,
                        mode=req.mode or "flash"
                    )
                    for chunk in stream:
                        full_response += chunk
                        yield f"data: {json.dumps({'chunk': chunk})}\n\n"
                except Exception as e:
                    logger.exception("Streaming LLM translation fallback failed")
                    full_response = "I encountered an error processing your request."
                    yield f"data: {json.dumps({'chunk': full_response})}\n\n"

        elif req.image_base64:
            try:
                vision_resp = call_llm_with_vision(
                    req.message, req.image_base64, req.image_mime or "image/jpeg",
                    system_prompt=system_override, file_context=file_context
                )
                full_response = vision_resp
                yield f"data: {json.dumps({'chunk': vision_resp})}\n\n"
            except Exception as e:
                logger.exception("Vision LLM failed")
                full_response = "I encountered an error processing the image."
                yield f"data: {json.dumps({'chunk': full_response})}\n\n"

        else:
            from llm import _detect_word_count_constraint
            word_target = _detect_word_count_constraint(req.message)
            if word_target is not None:
                try:
                    constrained_resp = call_llm_with_constraints(
                        req.message,
                        base_system_prompt=system_override,
                        custom_instructions=custom_instructions,
                        personality=personality,
                        max_tokens=1200,
                        mode=req.mode or "flash",
                    )
                    full_response = constrained_resp
                    yield f"data: {json.dumps({'chunk': constrained_resp})}\n\n"
                except Exception as e:
                    logger.exception("Constrained LLM failed")
                    full_response = "I encountered an error processing your request."
                    yield f"data: {json.dumps({'chunk': full_response})}\n\n"
            else:
                try:
                    stream = call_llm_streaming(
                        req.message,
                        custom_instructions=custom_instructions,
                        personality=personality,
                        system_prompt_override=system_override,
                        mode=req.mode or "flash"
                    )
                    for chunk in stream:
                        full_response += chunk
                        yield f"data: {json.dumps({'chunk': chunk})}\n\n"
                except Exception as e:
                    logger.exception("Streaming LLM failed")
                    full_response = "I encountered an error processing your request."
                    yield f"data: {json.dumps({'chunk': full_response})}\n\n"

        latency_ms = int((time.time() - start_time) * 1000)
        metrics.record_latency(latency_ms)
        if aug_meta.get("adaptation_used", False):
            metrics.inc("adaptation_used_count")
        message_id = str(uuid.uuid4())
        session_id = req.session_id or str(uuid.uuid4())
        if not req.session_id:
            db.create_session(session_id, user_id, title=req.message[:50])
        db.save_message(
            message_id=message_id,
            session_id=session_id,
            user_id=user_id,
            user_message=req.message,
            agent_response=full_response,
            intent=intent,
            sentiment=sentiment,
            frustration=frustration,
            latency_ms=latency_ms
        )
        memory.maybe_refresh_user_memory(user_id, session_id, req.message, full_response)
        try:
            adaptation.observe_interaction(
                user_id=user_id,
                user_message=req.message,
                agent_response=full_response,
                intent=intent,
                sentiment=sentiment.get("label") if isinstance(sentiment, dict) else str(sentiment),
                adaptation_used=aug_meta.get("adaptation_used", False),
            )
        except Exception as e:
            logger.warning(f"Adaptation observation failed: {e}")

        log_event(
            "chat_stream_completed",
            latency_ms=latency_ms,
            intent=intent,
            adaptation_used=aug_meta.get("adaptation_used", False),
            strategy=aug_meta.get("strategy")
        )

        yield f"data: {json.dumps({'done': True, 'message_id': message_id, 'session_id': session_id, 'metadata': {'intent': intent, 'sentiment': sentiment, 'frustration_score': frustration, 'searched': aug_meta.get('searched', False), 'deepl_used': aug_meta.get('deepl_used', False), 'strategy': aug_meta.get('strategy'), 'adaptation_used': aug_meta.get('adaptation_used', False)}})}\n\n"

    return StreamingResponse(generate(), media_type="text/event-stream")

# ─── Voice transcription ─────────────────────────────────────────────────
@app.post("/api/transcribe")
@limiter.limit("10/minute")
async def transcribe_audio(request: Request, audio: UploadFile = File(...)):
    require_user(request)
    GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
    if not GROQ_API_KEY:
        raise HTTPException(status_code=503, detail="GROQ_API_KEY not set")
    audio_bytes = await audio.read()
    if len(audio_bytes) < 1000:
        raise HTTPException(status_code=400, detail="Audio too short")
    if len(audio_bytes) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Audio file exceeds 10MB limit")
    content_type = audio.content_type or "audio/webm"
    ext = "webm" if "webm" in content_type else "ogg" if "ogg" in content_type else "m4a" if "m4a" in content_type else "webm"
    filename = f"recording.{ext}"
    try:
        import requests as _req
        response = _req.post(
            "https://api.groq.com/openai/v1/audio/transcriptions",
            headers={"Authorization": f"Bearer {GROQ_API_KEY}"},
            files={"file": (filename, audio_bytes, content_type)},
            data={"model": "whisper-large-v3", "response_format": "json", "language": "en"},
            timeout=30,
        )
        if response.status_code == 200:
            data = response.json()
            return {"text": data.get("text", "").strip()}
        else:
            raise HTTPException(status_code=response.status_code, detail=f"Groq Whisper error: {response.text[:200]}")
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Audio transcription failed")
        raise HTTPException(status_code=500, detail="Transcription failed")

# ─── Code execution ─────────────────────────────────────────────────────
class RunRequest(BaseModel):
    lang: str = "python"
    code: str
    stdin: str = ""

@app.post("/api/run")
@limiter.limit("10/minute")
def run_code(req: RunRequest, request: Request):
    require_user(request)
    result = powers.execute_code(req.lang, req.code, req.stdin)
    formatted = powers.format_execution_result(req.lang, req.code, result)
    return {"output": formatted, "raw": result}

# ─── Suggestions ──────────────────────────────────────────────────────
@app.post("/api/suggestions")
def get_suggestions(req: SuggestionRequest, request: Request):
    current_user = require_user(request)
    if req.user_id and req.user_id != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")
    past = db.get_similar_past_queries(current_user, req.current_input)
    return {"suggestions": past}

# ─── PPT ──────────────────────────────────────────────────────────────
@app.post("/api/generate-ppt")
@limiter.limit("5/minute")
async def generate_ppt(req: PPTRequest, request: Request):
    require_user(request)
    if not PPT_AVAILABLE:
        raise HTTPException(status_code=501, detail="python-pptx not installed")
    prompt = f"Generate exactly {req.slides} slides for a presentation about \"{req.topic}\". Return ONLY a JSON array of objects with keys: \"title\", \"content\". Example: [{{\"title\": \"Intro\", \"content\": \"...\"}}]"
    try:
        llm_resp = call_llm(prompt, mode="flash")
        start = llm_resp.find('[')
        end = llm_resp.rfind(']') + 1
        if start == -1 or end == 0:
            slides_data = [{"title": req.topic, "content": llm_resp[:200]}]
        else:
            slides_data = json.loads(llm_resp[start:end])
    except Exception:
        slides_data = [{"title": req.topic, "content": f"Presentation about {req.topic}"}]
    prs = Presentation()
    for slide in slides_data[:req.slides]:
        slide_layout = prs.slide_layouts[1]
        s = prs.slides.add_slide(slide_layout)
        s.shapes.title.text = slide.get("title", "Slide")
        s.placeholders[1].text = slide.get("content", "")
    ppt_bytes = BytesIO()
    prs.save(ppt_bytes)
    ppt_bytes.seek(0)
    safe_topic = re.sub(r"[^\w\-.]", "_", req.topic)
    return Response(
        content=ppt_bytes.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f"attachment; filename={safe_topic}.pptx"}
    )

# ─── Image Generation ─────────────────────────────────────────────────
@app.post("/api/generate-image")
@limiter.limit("5/minute")
async def generate_image(req: ImageRequest, request: Request):
    require_user(request)
    if not REPLICATE_AVAILABLE or not os.getenv("REPLICATE_API_KEY"):
        return {"image_url": "https://placekitten.com/512/512", "mock": True}
    try:
        output = replicate.run(
            "stability-ai/stable-diffusion:db21e45d3f7023abc2a46ee38a23973f6dce16bb082a930b0c49861f96d1e5bf",
            input={"prompt": req.prompt, "width": 512, "height": 512}
        )
        url = output[0] if isinstance(output, list) else output
        return {"image_url": url, "mock": False}
    except Exception as e:
        logger.exception("Image generation error")
        raise HTTPException(status_code=500, detail="Image generation failed")

# ─── Feedback ─────────────────────────────────────────────────────────
@app.post("/api/feedback")
@limiter.limit("30/minute")
def feedback(req: FeedbackRequest, request: Request):
    current_user = require_user(request)
    if req.user_id and req.user_id != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")
    user_id = current_user

    if req.session_id:
        sess = db.get_session(req.session_id)
        if not sess:
            raise HTTPException(status_code=404, detail="Session not found")
        if sess["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    if req.message_id:
        msg = db.get_message(req.message_id)
        if msg and msg["user_id"] != current_user:
            raise HTTPException(status_code=403, detail="Forbidden: message does not belong to you")

    db.save_feedback(user_id, req.session_id, req.message_id, req.helpful)
    if req.helpful:
        metrics.inc("feedback_positive")
    else:
        metrics.inc("feedback_negative")
        sol = db.get_last_agent_response(user_id, req.session_id)
        if sol:
            db.save_failed_solution(user_id, req.session_id, sol)
    try:
        adaptation.process_feedback(user_id, req.message_id, req.helpful)
    except Exception as e:
        logger.warning(f"Adaptation feedback processing failed: {e}")

    log_event("feedback_recorded", helpful=req.helpful)
    return {"status": "recorded"}

# ─── Adaptation Endpoints ─────────────────────────────────────────────
@app.get("/api/adaptation/{user_id}")
def get_user_adaptation(user_id: str, request: Request):
    require_same_user(request, user_id)
    profile = adaptation.get_adaptation_profile(user_id)
    context = adaptation.get_adaptation_context(user_id)
    strategy_stats = db.get_strategy_stats(user_id)
    preferred_strategies = adaptation.get_preferred_strategies(user_id)
    policy = adaptation.get_behavior_policy(user_id)
    metrics = adaptation.get_adaptation_metrics(user_id)
    return {
        "user_id": user_id,
        "profile": profile,
        "context": context,
        "preferred_strategies": preferred_strategies,
        "strategy_stats": strategy_stats,
        "policy": policy,
        "metrics": metrics,
    }

@app.delete("/api/adaptation/{user_id}")
def reset_user_adaptation(user_id: str, request: Request):
    require_same_user(request, user_id)
    adaptation.reset_adaptation_profile(user_id)
    log_event("adaptation_profile_reset", user_id=user_id)
    return {"status": "reset", "user_id": user_id, "message": "Adaptation profile reset to defaults"}

class AdaptationPatchRequest(BaseModel):
    preference: str
    value: Any
    confidence: Optional[Any] = 0.85

@app.patch("/api/adaptation/{user_id}")
def patch_user_adaptation(user_id: str, req: AdaptationPatchRequest, request: Request):
    require_same_user(request, user_id)
    try:
        conf = req.confidence if req.confidence is not None else 0.85
        updated = adaptation.set_manual_preference(user_id, req.preference, req.value, conf)
        log_event("adaptation_preference_patched", user_id=user_id, preference=req.preference)
        return {"status": "updated", "user_id": user_id, "profile": updated}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=400, detail="Invalid request parameters.")

# ─── Sessions ─────────────────────────────────────────────────────────
@app.get("/api/sessions/{user_id}")
def list_sessions(user_id: str, request: Request):
    require_same_user(request, user_id)
    return {"sessions": db.get_sessions_for_user(user_id)}

@app.get("/api/sessions/{user_id}/{session_id}/messages")
def session_messages(user_id: str, session_id: str, request: Request):
    require_same_user(request, user_id)
    sess = db.get_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found")
    if sess["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    return {"messages": db.get_session_messages(session_id)}

@app.delete("/api/sessions/{user_id}/{session_id}")
def delete_session(user_id: str, session_id: str, request: Request):
    require_same_user(request, user_id)
    sess = db.get_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found")
    if sess["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    deleted = db.delete_session(user_id, session_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Session not found")
    return {"status": "deleted"}

@app.put("/api/sessions/{user_id}/{session_id}/rename")
def rename_session(user_id: str, session_id: str, req: RenameRequest, request: Request):
    require_same_user(request, user_id)
    sess = db.get_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found")
    if sess["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    success = db.update_session_title(user_id, session_id, req.new_title.strip())
    if not success:
        raise HTTPException(status_code=404, detail="Session not found")
    return {"status": "updated"}

@app.post("/api/sessions/{user_id}/{session_id}/pin")
def pin_session(user_id: str, session_id: str, pinned: bool, request: Request):
    require_same_user(request, user_id)
    sess = db.get_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found")
    if sess["user_id"] != user_id:
        raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    db.set_pin(user_id, session_id, pinned)
    return {"status": "ok"}

@app.get("/api/search/{user_id}")
def search_sessions(user_id: str, q: str, request: Request):
    require_same_user(request, user_id)
    return {"sessions": db.search_sessions(user_id, q)}

# ─── User Preferences ──────────────────────────────────────────────────
@app.get("/api/user/preferences/{user_id}")
def get_preferences(user_id: str, request: Request):
    require_same_user(request, user_id)
    prefs = db.get_user_preferences(user_id)
    if not prefs:
        prefs = {"custom_instructions": "", "personality": "friendly", "theme": "light"}
    return prefs

@app.put("/api/user/preferences/{user_id}")
def update_preferences(user_id: str, req: PreferencesUpdate, request: Request):
    require_same_user(request, user_id)
    update_data = {k: v for k, v in req.dict().items() if v is not None}
    db.update_user_preferences(user_id, update_data)
    return {"status": "updated"}

# ─── Password Change ───────────────────────────────────────────────────
@app.put("/api/user/password/{user_id}")
@limiter.limit("5/minute")
def change_password(user_id: str, req: ChangePasswordRequest, request: Request):
    require_same_user(request, user_id)
    user = db.get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    if not db.verify_password(req.current_password, user["password_hash"], user["password_salt"]):
        raise HTTPException(status_code=401, detail="Current password is incorrect")
    if len(req.new_password) < 8:
        raise HTTPException(status_code=400, detail="New password must be at least 8 characters")
    db.update_password(user_id, req.new_password)
    log_event("password_changed", user_id=user_id)
    return {"status": "updated"}

# ─── Terms Acceptance ───────────────────────────────────────────────────
@app.get("/api/user/terms/{user_id}")
def get_terms_status(user_id: str, request: Request):
    require_same_user(request, user_id)
    user = db.get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    return {
        "terms_accepted": bool(user.get("terms_accepted", 0)),
        "terms_accepted_date": user.get("terms_accepted_date")
    }

@app.put("/api/user/terms/{user_id}")
def update_terms_acceptance(user_id: str, req: TermsAcceptRequest, request: Request):
    require_same_user(request, user_id)
    user = db.get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    db.update_terms_acceptance(user_id, req.terms_accepted)
    return {
        "terms_accepted": req.terms_accepted,
        "terms_accepted_date": datetime.now(timezone.utc).isoformat() if req.terms_accepted else None
    }

# ─── Document Upload ──────────────────────────────────────────────────
UPLOAD_DIR = config.UPLOAD_DIR
os.makedirs(UPLOAD_DIR, exist_ok=True)
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".csv", ".tsv", ".txt", ".md", ".log"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB

@app.post("/api/upload-document")
async def upload_document(
    request: Request,
    user_id: str = Form(...),
    file: UploadFile = File(...)
):
    current_user = require_user(request)
    if user_id != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")

    filename = os.path.basename(file.filename)
    file_ext = os.path.splitext(filename)[1].lower()
    if file_ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail=f"File extension '{file_ext}' is not supported")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File size exceeds maximum allowed 10MB")

    doc_id = str(uuid.uuid4())
    safe_filename = f"{doc_id}{file_ext}"
    file_path = os.path.join(UPLOAD_DIR, safe_filename)
    with open(file_path, "wb") as f:
        f.write(content)

    text = ""
    if file_ext == ".pdf":
        try:
            import PyPDF2
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text
        except Exception as e:
            text = f"[PDF extraction error: {str(e)}]"
    elif file_ext == ".docx":
        try:
            import docx
            doc = docx.Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            text = f"[DOCX extraction error: {str(e)}]"
    elif file_ext in [".xlsx", ".xls"]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values=True):
                    text += " ".join(str(cell) for cell in row if cell) + "\n"
        except Exception as e:
            text = f"[Excel extraction error: {str(e)}]"
    elif file_ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            text = f"[PPTX extraction error: {str(e)}]"
    elif file_ext in [".csv", ".tsv"]:
        try:
            import csv as csv_module
            delimiter = "\t" if file_ext == ".tsv" else ","
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv_module.reader(f, delimiter=delimiter)
                rows = list(reader)
            if rows:
                headers = rows[0]
                data_rows = rows[1:]
                row_count = len(data_rows)
                col_count = len(headers)
                text = f"[CSV DATA — {row_count} rows × {col_count} columns]\n"
                text += f"Columns: {', '.join(headers)}\n\n"
                preview_rows = data_rows[:200]
                text += "\n".join([",".join(r) for r in [headers] + preview_rows])
                if row_count > 200:
                    text += f"\n\n[... {row_count - 200} more rows not shown. Total rows: {row_count}]"
        except Exception as e:
            text = f"[CSV extraction error: {str(e)}]"
    elif file_ext in [".txt", ".md", ".log"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()[:50000]
        except Exception as e:
            text = f"[Text extraction error: {str(e)}]"

    db.save_document(doc_id, user_id, filename, file_path, text)
    log_event("document_uploaded", user_id=user_id, doc_id=doc_id)
    return {
        "doc_id": doc_id,
        "filename": filename,
        "text_preview": text[:300]
    }

# ─── Share ────────────────────────────────────────────────────────────
@app.get("/api/share/{session_id}")
def get_share_link(session_id: str, request: Request):
    current_user = require_user(request)
    sess = db.get_session(session_id)
    if not sess:
        raise HTTPException(status_code=404, detail="Session not found")
    if sess["user_id"] != current_user:
        raise HTTPException(status_code=403, detail="Forbidden: session does not belong to you")
    token = db.get_or_create_share_token(session_id, current_user)
    return {"share_url": f"/share/{token}"}

@app.get("/share/{token}")
async def view_shared_chat(token: str):
    data = db.get_shared_session_data(token)
    if not data:
        raise HTTPException(status_code=404, detail="Share link invalid or expired")
    html = f"""
    <!DOCTYPE html>
    <html>
    <head><meta charset="UTF-8"><title>Shared Chat - Ava AI</title>
    <style>body{{font-family:Inter,sans-serif;background:#f5f4ff;padding:20px;max-width:760px;margin:40px auto;color:#1d1230;}}
    .bubble{{padding:12px 18px;border-radius:18px;margin-bottom:12px;line-height:1.6;}}
    .user{{background:#e3cdfb;border-radius:18px 18px 4px 18px;}}
    .ava{{background:#ffffff;border:1px solid #ece2fb;border-radius:18px 18px 18px 4px;}}
    .meta{{font-size:0.8rem;color:#9a8db8;margin-top:6px;}}
    </style>
    </head>
    <body>
    <h1 style="font-family:Sora,sans-serif;">Shared Conversation</h1>
    <div id="messages">
    """
    for msg in data:
        html += f"""
        <div class="bubble user"><strong>User:</strong> {msg['user_message']}</div>
        <div class="bubble ava"><strong>Ava:</strong> {msg['agent_response']}</div>
        <div class="meta">Sentiment: {msg.get('sentiment_label','N/A')} • Intent: {msg.get('intent','N/A')}</div>
        """
    html += "</div></body></html>"
    return HTMLResponse(content=html)

@app.get("/api/share-data")
def share_data(token: str):
    data = db.get_shared_session_data(token)
    if not data:
        raise HTTPException(status_code=404, detail="Invalid share link")
    return data

# ─── Export ────────────────────────────────────────────────────────────
@app.get("/api/export/{user_id}")
def export_data(user_id: str, request: Request):
    require_same_user(request, user_id)
    user = db.get_user_by_id(user_id)
    sessions = db.get_sessions_for_user(user_id)
    messages = db.get_all_messages(user_id)
    feedback_data = db.get_feedback(user_id)
    docs = db.get_user_documents(user_id)
    data = {
        "user": user,
        "sessions": sessions,
        "messages": messages,
        "feedback": feedback_data,
        "documents": docs
    }
    json_str = json.dumps(data, indent=2, default=str)
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        zf.writestr("ava_export.json", json_str)
    zip_buffer.seek(0)
    return Response(
        content=zip_buffer.getvalue(),
        media_type="application/zip",
        headers={"Content-Disposition": "attachment; filename=ava_export.zip"}
    )

# ─── Web Search ────────────────────────────────────────────────────────
@app.get("/api/web-search")
def web_search(query: str, request: Request):
    require_user(request)
    results = search.web_search(query)
    return {"results": results}

# ─── Memory ───────────────────────────────────────────────────────────
@app.get("/api/memory/{user_id}")
def get_memory(user_id: str, request: Request):
    require_same_user(request, user_id)
    stats = db.get_user_stats(user_id)
    recent = db.get_recent_messages(user_id, limit=5)
    user_mem = db.get_user_memory(user_id)
    return {"stats": stats, "recent": recent, "remembered_facts": user_mem["memory_text"]}

@app.delete("/api/memory/{user_id}")
def clear_memory(user_id: str, request: Request):
    require_same_user(request, user_id)
    db.delete_user_data(user_id)
    return {"message": "Data cleared"}

# ─── Cross-session user memory ──────────────────────────────────────
@app.get("/api/user-memory/{user_id}")
def get_user_memory(user_id: str, request: Request):
    require_same_user(request, user_id)
    return db.get_user_memory(user_id)

@app.delete("/api/user-memory/{user_id}")
def clear_user_memory(user_id: str, request: Request):
    require_same_user(request, user_id)
    db.clear_user_memory(user_id)
    return {"status": "cleared"}

# ─── Reviews ──────────────────────────────────────────────────────────
REVIEW_IMAGE_DIR = config.REVIEW_IMAGE_DIR
os.makedirs(REVIEW_IMAGE_DIR, exist_ok=True)
ALLOWED_IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
MAX_IMAGE_BYTES = 5 * 1024 * 1024

class ReviewSubmitRequest(BaseModel):
    rating: int = Field(..., ge=1, le=5)
    comment: Optional[str] = Field("", max_length=2000)
    user_id: Optional[str] = None

@app.post("/api/reviews")
async def submit_review(payload: ReviewSubmitRequest, request: Request):
    try:
        reviewer_name = "Guest"
        if payload.user_id:
            current_user = request.session.get("user_id")
            if current_user and current_user != payload.user_id:
                raise HTTPException(status_code=403, detail="Forbidden: user_id mismatch")
            user = db.get_user_by_id(payload.user_id)
            if user:
                reviewer_name = user["name"]

        review_id = str(uuid.uuid4())
        db.create_review(
            review_id=review_id,
            reviewer_name=reviewer_name,
            reviewer_title=None,
            rating=payload.rating,
            review_text=payload.comment,
            image_path=None
        )
        log_event("review_submitted", rating=payload.rating)
        return {"review_id": review_id, "status": "submitted"}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Review submission error")
        raise HTTPException(status_code=500, detail="Review submission failed")

def _serialize_review_for_public(r: dict) -> dict:
    return {
        "rating": r["rating"],
        "comment": r["review_text"] or "",
        "user_id": r["reviewer_name"] or "Guest",
        "created_at": r["created_at"]
    }

@app.get("/api/reviews")
def list_public_reviews():
    reviews = db.get_public_reviews()
    return [_serialize_review_for_public(r) for r in reviews]

@app.get("/api/reviews/public")
def public_reviews_alias():
    return list_public_reviews()

# ─── Admin review endpoints ─────────────────────────────────────────
@app.get("/api/admin/reviews")
def admin_list_reviews(request: Request):
    require_admin(request)
    reviews = db.get_all_reviews()
    result = []
    for r in reviews:
        result.append({
            "id": r["review_id"],
            "rating": r["rating"],
            "comment": r["review_text"] or "",
            "user_id": r["reviewer_name"] or "Guest",
            "created_at": r["created_at"],
            "status": r["status"],
            "featured": bool(r["featured"]),
            "reviewer_name": r["reviewer_name"],
            "review_text": r["review_text"] or ""
        })
    return {"reviews": result}

class ReviewStatusUpdate(BaseModel):
    status: str

@app.patch("/api/admin/reviews/{review_id}/status")
def admin_update_review_status(review_id: str, req: ReviewStatusUpdate, request: Request):
    require_admin(request)
    if req.status not in ("approved", "hidden", "pending"):
        raise HTTPException(status_code=400, detail="Invalid status")
    if not db.get_review(review_id):
        raise HTTPException(status_code=404, detail="Review not found")
    db.update_review_status(review_id, req.status)
    return {"status": "updated"}

class ReviewFeaturedUpdate(BaseModel):
    featured: bool

@app.patch("/api/admin/reviews/{review_id}/featured")
def admin_update_review_featured(review_id: str, req: ReviewFeaturedUpdate, request: Request):
    require_admin(request)
    if not db.get_review(review_id):
        raise HTTPException(status_code=404, detail="Review not found")
    db.set_review_featured(review_id, req.featured)
    return {"status": "updated"}

@app.delete("/api/admin/reviews/{review_id}")
def admin_delete_review(review_id: str, request: Request):
    require_admin(request)
    review = db.delete_review(review_id)
    if not review:
        raise HTTPException(status_code=404, detail="Review not found")
    if review.get("image_path") and os.path.exists(review["image_path"]):
        try:
            os.remove(review["image_path"])
        except OSError:
            pass
    return {"status": "deleted"}

# ─── Analytics ────────────────────────────────────────────────────────
@app.get("/api/analytics")
def analytics(request: Request):
    require_admin(request)
    return db.get_global_analytics()

@app.get("/api/analytics/{user_id}")
def user_analytics(user_id: str, request: Request):
    require_same_user(request, user_id)
    return db.get_user_stats(user_id)

# ─── Static Files & Admin Page ───────────────────────────────────────
STATIC_DIR = (
    os.path.join(os.path.dirname(__file__), "static")
    if os.path.isdir(os.path.join(os.path.dirname(__file__), "static"))
    else os.path.dirname(__file__)
)

# ─── Cache-busting route for admin.html (protected) ──────────────────
@app.get("/admin.html")
async def admin_page(request: Request):
    require_admin(request)
    return FileResponse(
        os.path.join(STATIC_DIR, "admin.html"),
        headers={"Cache-Control": "no-cache, no-store, must-revalidate"}
    )

app.mount("/", StaticFiles(directory=STATIC_DIR, html=True), name="static")